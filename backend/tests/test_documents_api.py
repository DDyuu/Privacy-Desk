def test_full_document_flow(unlocked_client):
    client = unlocked_client
    content = "联系人张三，电话13800138000，身份证11010519491231002X。".encode("utf-8")
    upload = client.post(
        "/api/v1/documents",
        files={"file": ("sample.txt", content, "text/plain")},
    )
    assert upload.status_code == 200
    doc_id = upload.json()["id"]

    listing = client.get("/api/v1/documents")
    assert listing.status_code == 200
    assert listing.json()[0]["id"] == doc_id

    detected = client.post(
        f"/api/v1/documents/{doc_id}/detect", json={"use_llm": False}
    )
    assert detected.status_code == 200
    codes = {item["type_code"] for item in detected.json()["spans"]}
    assert {"PHONE_CN", "ID_CARD_CN"} <= codes
    assert detected.json()["mode"] == "rules"

    masked = client.post(
        f"/api/v1/documents/{doc_id}/mask", json={"actions": {}, "verify": True}
    )
    assert masked.status_code == 200
    body = masked.json()
    assert "13800138000" not in body["masked_text"]
    assert "11010519491231002X" not in body["masked_text"]
    assert body["verify"]["passed"] is True
    assert body["stats"]["detected"] >= 2

    compare = client.get(f"/api/v1/documents/{doc_id}/compare")
    assert "13800138000" in compare.json()["original"]
    assert "13800138000" not in compare.json()["masked"]

    no_password = client.get(
        f"/api/v1/documents/{doc_id}/export", params={"kind": "report", "fmt": "json"}
    )
    assert no_password.status_code == 401

    wrong_password = client.get(
        f"/api/v1/documents/{doc_id}/export",
        params={"kind": "report", "fmt": "json", "password": "wrong-pass"},
    )
    assert wrong_password.status_code == 403

    report = client.get(
        f"/api/v1/documents/{doc_id}/export",
        params={"kind": "report", "fmt": "json", "password": "password123"},
    )
    assert report.status_code == 200
    assert report.json()["summary"]["detected"] >= 2

    exported = client.get(
        f"/api/v1/documents/{doc_id}/export", params={"kind": "masked", "fmt": "txt"}
    )
    assert exported.status_code == 200
    assert "13800138000" not in exported.text


def test_edit_content_and_reject_unsupported(unlocked_client):
    client = unlocked_client
    upload = client.post(
        "/api/v1/documents",
        files={"file": ("a.txt", "电话13800138000".encode("utf-8"), "text/plain")},
    )
    doc_id = upload.json()["id"]
    edited = client.put(
        f"/api/v1/documents/{doc_id}/content",
        json={"text": "新的正文 13800138000"},
    )
    assert edited.status_code == 200
    assert "新的正文" in client.get(f"/api/v1/documents/{doc_id}").json()["text"]

    bad = client.post(
        "/api/v1/documents",
        files={"file": ("a.exe", b"MZ", "application/octet-stream")},
    )
    assert bad.status_code == 400


def test_docx_export_uses_applied_mapping(unlocked_client, tmp_path):
    from docx import Document
    from io import BytesIO

    doc = Document()
    doc.add_paragraph("姓名张三")
    doc.add_paragraph("电话13800138000")
    buffer = BytesIO()
    doc.save(buffer)
    upload = unlocked_client.post(
        "/api/v1/documents",
        files={
            "file": (
                "sample.docx",
                buffer.getvalue(),
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        },
    )
    assert upload.status_code == 200
    doc_id = upload.json()["id"]
    unlocked_client.post(f"/api/v1/documents/{doc_id}/detect", json={"use_llm": False})
    unlocked_client.post(
        f"/api/v1/documents/{doc_id}/mask",
        json={"actions": {"PHONE_CN": {"action": "replace"}}, "verify": True},
    )
    exported = unlocked_client.get(
        f"/api/v1/documents/{doc_id}/export", params={"kind": "masked", "fmt": "docx"}
    )
    assert exported.status_code == 200
    reloaded = Document(BytesIO(exported.content))
    assert "13800138000" not in reloaded.paragraphs[1].text
    assert "电话" in reloaded.paragraphs[1].text
