from docx import Document
from docx.oxml import OxmlElement
import pytest

from app.core.ingest.docx_loader import load_docx
from app.core.ingest.registry import export_document, load_document


def test_docx_roundtrip_preserves_table(tmp_path):
    source = tmp_path / "a.docx"
    doc = Document()
    doc.add_paragraph("联系人电话 13800138000")
    table = doc.add_table(rows=1, cols=1)
    table.cell(0, 0).text = "身份证 11010519491231002X"
    doc.save(source)

    model = load_docx(source)
    assert "13800138000" in model.text
    assert "11010519491231002X" in model.text

    masked = model.text.replace("13800138000", "138****8000").replace(
        "11010519491231002X", "**************002X"
    )
    target = tmp_path / "out.docx"
    export_document(model, masked, target, "docx")

    reloaded = Document(target)
    assert reloaded.paragraphs[0].text == "联系人电话 138****8000"
    assert reloaded.tables[0].cell(0, 0).text == "身份证 **************002X"


def test_load_document_dispatches_docx(tmp_path):
    source = tmp_path / "b.docx"
    doc = Document()
    doc.add_paragraph("hello")
    doc.save(source)
    assert load_document(source, "docx").fmt == "docx"


def test_docx_export_remaps_offsets_for_longer_replacements(tmp_path):
    source = tmp_path / "long.docx"
    doc = Document()
    doc.add_paragraph("姓名张三")
    doc.add_paragraph("电话13800138000")
    doc.save(source)

    model = load_docx(source)
    text = model.text
    start = text.index("张三")
    masked = text.replace("张三", "[姓名_1]")
    applied = [
        {
            "start": start,
            "end": start + 2,
            "type_code": "NAME_CN",
            "action": "replace",
            "replacement": "[姓名_1]",
        }
    ]

    target = tmp_path / "long_out.docx"
    export_document(model, masked, target, "docx", applied=applied)

    reloaded = Document(target)
    assert reloaded.paragraphs[0].text == "姓名[姓名_1]"
    assert reloaded.paragraphs[1].text == "电话13800138000"


def test_docx_export_clears_hyperlink_text(tmp_path):
    source = tmp_path / "link.docx"
    doc = Document()
    paragraph = doc.add_paragraph("联系电话：")
    hyperlink = OxmlElement("w:hyperlink")
    run = OxmlElement("w:r")
    text = OxmlElement("w:t")
    text.text = "13800138000"
    run.append(text)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)
    doc.save(source)

    model = load_docx(source)
    assert "13800138000" in model.text
    masked = model.text.replace("13800138000", "*******8000")
    start = model.text.index("13800138000")
    applied = [
        {
            "start": start,
            "end": start + 11,
            "type_code": "PHONE_CN",
            "action": "mask",
            "replacement": "*******8000",
        }
    ]
    target = tmp_path / "link_out.docx"
    export_document(model, masked, target, "docx", applied=applied)
    reloaded = Document(target)
    assert "13800138000" not in reloaded.paragraphs[0].text
    assert "*******8000" in reloaded.paragraphs[0].text


def test_docx_merged_cell_not_duplicated(tmp_path):
    source = tmp_path / "merge.docx"
    doc = Document()
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "身份证11010519491231002X"
    table.cell(0, 0).merge(table.cell(1, 1))
    doc.save(source)
    model = load_docx(source)
    assert model.text.count("11010519491231002X") == 1


def test_docx_nested_table_is_scanned(tmp_path):
    source = tmp_path / "nested.docx"
    doc = Document()
    outer = doc.add_table(rows=1, cols=1)
    cell = outer.cell(0, 0)
    cell.text = "外层"
    nested = cell.add_table(rows=1, cols=1)
    nested.cell(0, 0).text = "身份证11010519491231002X"
    doc.save(source)
    model = load_docx(source)
    assert "11010519491231002X" in model.text


def test_docx_export_rejects_missing_applied_when_length_changes(tmp_path):
    source = tmp_path / "guard.docx"
    doc = Document()
    doc.add_paragraph("姓名张三")
    doc.save(source)
    model = load_docx(source)
    masked = model.text.replace("张三", "[姓名_1]")
    with pytest.raises(ValueError, match="applied"):
        export_document(model, masked, tmp_path / "guard_out.docx", "docx")
